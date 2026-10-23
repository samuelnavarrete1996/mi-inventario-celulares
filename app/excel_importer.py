import io
import re
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Tuple, Optional
from app import crud
from app.models import EquipoCreate

def clean_imei_val(val: Any) -> str:
    """Convierte de forma segura cualquier formato de IMEI (incluso notación científica de Excel) a string exacto."""
    if val is None:
        return ""
    if isinstance(val, float):
        # Evitar notación científica ej. 3.56789123456781e+14 -> 356789123456781
        val_str = f"{val:.0f}"
    else:
        val_str = str(val).strip()

    # Si viene con notación científica en texto ej: '3.56789E+14'
    if 'e+' in val_str.lower():
        try:
            val_str = f"{float(val_str):.0f}"
        except Exception:
            pass

    return val_str.replace(" ", "").replace("-", "").replace(".0", "")

def extraer_capacidad_de_modelo(modelo: str) -> int:
    """Intenta deducir la capacidad en GB si no está en una columna separada."""
    match = re.search(r'(\d+)\s*(?:gb|g\b|tb)', modelo.lower())
    if match:
        num = int(match.group(1))
        if num == 1: # 1 TB
            return 1024
        return num
    return 128

def normalizar_grado(grado_val: Any) -> str:
    g = str(grado_val or '').strip().upper()
    if g.startswith('A') or 'EXCELENTE' in g:
        return 'A'
    elif g.startswith('B') or 'REGULAR' in g or 'BUENO' in g:
        return 'B'
    elif g.startswith('C') or 'MALTRATADO' in g or 'DETALLES' in g:
        return 'C'
    return 'A' # Por defecto Grado A

def mapear_fila_a_equipo(row_dict: Dict[str, Any], fila_num: int) -> Tuple[Optional[EquipoCreate], Optional[str]]:
    """Convierte un diccionario de fila de Excel a un objeto EquipoCreate validado."""
    # Normalizar llaves
    norm = {}
    for k, v in row_dict.items():
        if k is not None:
            clean_k = str(k).strip().lower().replace(" ", "").replace("_", "").replace("-", "")
            norm[clean_k] = v

    # Buscar Modelo
    modelo = norm.get('modelo') or norm.get('nombre') or norm.get('equipo') or norm.get('telefono') or norm.get('dispositivo')
    if not modelo or not str(modelo).strip():
        return None, f"Fila {fila_num}: Falta el 'Modelo' del teléfono."

    modelo_str = str(modelo).strip()

    # Buscar Color
    color = norm.get('color') or norm.get('tono') or 'Negro'
    color_str = str(color).strip()

    # Buscar IMEI
    imei_raw = norm.get('imei') or norm.get('serie') or norm.get('serial') or norm.get('numeroserie') or norm.get('sn')
    imei_str = clean_imei_val(imei_raw)
    if not imei_str or len(imei_str) < 8:
        return None, f"Fila {fila_num}: IMEI inválido o ausente ({imei_raw})."

    # Buscar Grado
    grado_raw = norm.get('grado') or norm.get('condicion') or norm.get('calidad') or norm.get('estado')
    grado_str = normalizar_grado(grado_raw)

    # Buscar Costo
    costo_raw = norm.get('costo') or norm.get('costocompra') or norm.get('preciocompra') or norm.get('compra')
    try:
        costo_float = float(str(costo_raw).replace("$", "").replace(",", "").strip())
        if costo_float <= 0:
            raise ValueError()
    except Exception:
        return None, f"Fila {fila_num}: Costo de compra inválido ({costo_raw}). Debe ser un número mayor a 0."

    # Buscar Capacidad (opcional, si falta se deduce de modelo o 128)
    cap_raw = norm.get('capacidad') or norm.get('capacidadgb') or norm.get('gb') or norm.get('almacenamiento')
    if cap_raw:
        try:
            cap_int = int(re.sub(r'[^\d]', '', str(cap_raw)))
            if cap_int == 1:
                cap_int = 1024
        except Exception:
            cap_int = extraer_capacidad_de_modelo(modelo_str)
    else:
        cap_int = extraer_capacidad_de_modelo(modelo_str)

    # Buscar Precio Venta (opcional, si falta se sugiere costo + 30% de margen)
    precio_raw = norm.get('precioventa') or norm.get('precio') or norm.get('venta') or norm.get('pvp')
    if precio_raw:
        try:
            precio_float = float(str(precio_raw).replace("$", "").replace(",", "").strip())
        except Exception:
            precio_float = round(costo_float * 1.30, 2)
    else:
        precio_float = round(costo_float * 1.30, 2)

    try:
        equipo = EquipoCreate(
            modelo=modelo_str,
            capacidad_gb=cap_int,
            color=color_str,
            imei=imei_str,
            grado=grado_str,
            costo_compra=costo_float,
            precio_venta=precio_float
        )
        return equipo, None
    except Exception as e:
        return None, f"Fila {fila_num}: Error de validación: {str(e)}"

def procesar_archivo_excel_o_csv(contenido_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Lee un archivo Excel (.xlsx) o CSV en memoria y retorna los registros válidos y errores."""
    filename_lower = filename.lower()
    filas_dict = []

    if filename_lower.endswith('.csv'):
        # Intentar decodificar como utf-8 o latin-1
        texto = None
        for enc in ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252'):
            try:
                texto = contenido_bytes.decode(enc)
                break
            except Exception:
                continue

        if not texto:
            raise ValueError("No se pudo leer la codificación del archivo CSV.")

        reader = csv.DictReader(io.StringIO(texto))
        for row in reader:
            filas_dict.append(row)

    elif filename_lower.endswith(('.xlsx', '.xlsm')):
        wb = openpyxl.load_workbook(io.BytesIO(contenido_bytes), data_only=True)
        sheet = wb.active

        headers = []
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            headers.append(str(val).strip() if val is not None else f"col_{col_idx}")

        for row_idx in range(2, sheet.max_row + 1):
            row_data = {}
            has_data = False
            for col_idx, h in enumerate(headers, start=1):
                val = sheet.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip() != "":
                    has_data = True
                row_data[h] = val
            if has_data:
                filas_dict.append(row_data)
    else:
        raise ValueError("Formato no compatible. Por favor sube un archivo Excel (.xlsx) o CSV (.csv).")

    if not filas_dict:
        raise ValueError("El archivo está vacío o no contiene filas de datos debajo del encabezado.")

    equipos_validos: List[EquipoCreate] = []
    errores: List[str] = []
    imeis_vistos = set()

    for idx, r in enumerate(filas_dict, start=2):
        eq, err = mapear_fila_a_equipo(r, idx)
        if err:
            errores.append(err)
            continue

        # Verificar si el IMEI está repetido dentro del mismo archivo
        if eq.imei in imeis_vistos:
            errores.append(f"Fila {idx}: El IMEI {eq.imei} está duplicado dentro de este mismo archivo.")
            continue

        # Verificar si ya existe en la base de datos
        existente = crud.get_equipo_by_imei(eq.imei)
        if existente:
            errores.append(f"Fila {idx}: El IMEI {eq.imei} ya existe registrado en el inventario ({existente['modelo']}).")
            continue

        imeis_vistos.add(eq.imei)
        equipos_validos.append(eq)

    return {
        "total_filas": len(filas_dict),
        "validos": equipos_validos,
        "errores": errores
    }

def guardar_equipos_importados(
    equipos: List[EquipoCreate],
    tenant_id: Optional[int] = None,
    sucursal_id: Optional[int] = None
) -> List[Dict[str, Any]]:
    """Inserta en lote los equipos validados dentro del tenant del usuario."""
    guardados = []
    for eq in equipos:
        nuevo = crud.create_equipo(eq, tenant_id=tenant_id, sucursal_id=sucursal_id)
        guardados.append(nuevo)
    return guardados

def generar_plantilla_excel_bytes() -> bytes:
    """Genera un archivo Excel (.xlsx) profesional con estilos y datos de ejemplo para descargar."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario Celulares"

    # Encabezados
    headers = ["modelo", "capacidad_gb", "color", "imei", "grado", "costo_compra", "precio_venta"]
    ws.append(headers)

    # Filas de ejemplo
    ejemplos = [
        ["iPhone 14", 128, "Negro Medianoche", "356789123456790", "A", 9500.0, 13000.0],
        ["Samsung Galaxy S23", 256, "Verde", "356789123456791", "A", 9000.0, 12500.0],
        ["Xiaomi Redmi Note 12", 128, "Azul", "356789123456792", "B", 2800.0, 3900.0],
        ["iPhone 11", 64, "Blanco", "356789123456793", "C", 3500.0, 4800.0],
    ]

    for row in ejemplos:
        ws.append(row)

    # Estilos
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Azul
    alignment_center = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center

    # Ajustar ancho de columnas
    anchos = [26, 16, 20, 22, 12, 16, 16]
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = ancho

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def generar_plantilla_celulares_csv() -> bytes:
    """Genera un archivo CSV profesional optimizado al 100% para Google Sheets y Excel (UTF-8 con BOM)."""
    output = io.StringIO()
    writer = csv.writer(output, lineterminator='\r\n')
    
    headers = ["modelo", "capacidad_gb", "color", "imei", "grado", "costo_compra", "precio_venta"]
    writer.writerow(headers)
    
    ejemplos = [
        ["iPhone 14", 128, "Negro Medianoche", "356789123456790", "A", 9500.0, 13000.0],
        ["Samsung Galaxy S23", 256, "Verde", "356789123456791", "A", 9000.0, 12500.0],
        ["Xiaomi Redmi Note 12", 128, "Azul", "356789123456792", "B", 2800.0, 3900.0],
        ["iPhone 11", 64, "Blanco", "356789123456793", "C", 3500.0, 4800.0],
    ]
    for row in ejemplos:
        writer.writerow(row)
        
    return output.getvalue().encode('utf-8-sig')


def generar_plantilla_stock_accesorios_excel(
    solo_encabezados: bool = True,
    tenant_id: Optional[int] = None
) -> bytes:
    """Genera un archivo Excel (.xlsx) con la plantilla de accesorios. Por defecto solo incluye la fila superior de encabezados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Accesorios"

    # Encabezados exactos solicitados
    headers = [
        "ID",
        "Nombre del Accesorio",
        "Color",
        "Categoría",
        "Stock Actual",
        "NUEVO STOCK",
        "Menudeo",
        "Mayoreo",
        "Caja"
    ]
    ws.append(headers)

    accesorios_fisicos = []
    if not solo_encabezados:
        from app import crud
        accesorios = crud.get_accesorios(tenant_id=tenant_id)
        accesorios_fisicos = [a for a in accesorios if not a.get("es_servicio")]
        for a in accesorios_fisicos:
            ws.append([
                a["id"],
                a["nombre"],
                a.get("color") or "",
                a["categoria"],
                int(a.get("stock", 0) or 0),
                int(a.get("stock", 0) or 0),
                float(a.get("precio_venta") or 0.0),
                float(a.get("precio_20_pzs") or a.get("precio_mayoreo") or a.get("precio_venta") or 0.0),
                float(a.get("precio_caja") or a.get("costo_compra") or 0.0)
            ])

    # Estilos de encabezado
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Verde Teal
    nuevo_stock_fill = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid") # Verde Esmeralda Destacado
    alignment_center = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = nuevo_stock_fill if col_idx == 6 else header_fill
        cell.alignment = alignment_center

    # Resaltar la columna NUEVO STOCK en las filas de datos con fondo suave si hay datos
    editable_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    for r in range(2, len(accesorios_fisicos) + 2):
        cell_stock = ws.cell(row=r, column=6)
        cell_stock.fill = editable_fill
        cell_stock.alignment = alignment_center
        cell_stock.font = Font(name="Calibri", size=11, bold=True)

    # Anchos de columna
    anchos = [8, 44, 18, 22, 14, 18, 16, 16, 16]
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = ancho

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def generar_plantilla_stock_accesorios_csv(
    solo_encabezados: bool = True,
    tenant_id: Optional[int] = None
) -> bytes:
    """Genera un archivo CSV optimizado al 100% para Google Sheets con la plantilla de accesorios."""
    output = io.StringIO()
    writer = csv.writer(output, lineterminator='\r\n')

    # Encabezados exactos solicitados
    headers = [
        "ID",
        "Nombre del Accesorio",
        "Color",
        "Categoría",
        "Stock Actual",
        "NUEVO STOCK",
        "Menudeo",
        "Mayoreo",
        "Caja"
    ]
    writer.writerow(headers)

    if not solo_encabezados:
        from app import crud
        accesorios = crud.get_accesorios(tenant_id=tenant_id)
        accesorios_fisicos = [a for a in accesorios if not a.get("es_servicio")]
        for a in accesorios_fisicos:
            writer.writerow([
                a["id"],
                a["nombre"],
                a.get("color") or "",
                a["categoria"],
                int(a.get("stock", 0) or 0),
                int(a.get("stock", 0) or 0),
                float(a.get("precio_venta") or 0.0),
                float(a.get("precio_20_pzs") or a.get("precio_mayoreo") or a.get("precio_venta") or 0.0),
                float(a.get("precio_caja") or a.get("costo_compra") or 0.0)
            ])

    return output.getvalue().encode('utf-8-sig')


def importar_stock_accesorios_excel(
    contenido_bytes: bytes,
    filename: str = "",
    usuario: str = "Admin",
    tenant_id: Optional[int] = None
) -> Dict[str, Any]:
    """Procesa el archivo de stock de accesorios (compatible con Google Sheets .csv y Excel .xlsx) y actualiza masivamente registrando Kardex y precios si aplican."""
    from app import crud

    items_to_update = []
    items_to_create = []
    errores = []

    filename_lower = filename.lower()
    is_csv = filename_lower.endswith('.csv')
    if not is_csv and not contenido_bytes.startswith(b'PK'):
        is_csv = True

    if is_csv:
        texto = None
        for enc in ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252'):
            try:
                texto = contenido_bytes.decode(enc)
                break
            except Exception:
                continue

        if not texto:
            raise ValueError("No se pudo leer la codificación del archivo CSV de Google Sheets.")

        reader = csv.reader(io.StringIO(texto))
        filas = list(reader)
        if not filas:
            raise ValueError("El archivo está vacío.")

        headers = [str(c).strip().lower().replace(" ", "").replace("_", "") for c in filas[0]]
        id_col = None
        stock_col = None
        nom_col = None
        col_col = None
        cat_col = None
        men_col = None
        may_col = None
        caj_col = None

        for idx, h in enumerate(headers):
            if h in ("id", "codigo", "clave"):
                id_col = idx
            elif h in ("nuevostock", "stock", "stockfisico", "piezas", "cantidad"):
                stock_col = idx
            elif h in ("nombredelaccesorio", "nombre", "accesorio", "descripcion", "item"):
                nom_col = idx
            elif h in ("color", "variante"):
                col_col = idx
            elif h in ("categoria", "familia", "tipo"):
                cat_col = idx
            elif h in ("menudeo", "preciomenudeo", "precioventa", "precio"):
                men_col = idx
            elif h in ("mayoreo", "preciomayoreo"):
                may_col = idx
            elif h in ("caja", "preciocaja"):
                caj_col = idx

        if id_col is None and stock_col is None and nom_col is None:
            raise ValueError("El archivo debe contener al menos las columnas 'ID' y 'NUEVO STOCK' o 'Nombre del Accesorio'.")

        for row_idx, row in enumerate(filas[1:], start=2):
            if not row:
                continue
            acc_id = None
            if id_col is not None and len(row) > id_col and str(row[id_col]).strip():
                try:
                    raw_id = int(float(str(row[id_col]).strip()))
                    if raw_id > 0:
                        acc_id = raw_id
                except Exception:
                    acc_id = None

            val_stock = None
            if stock_col is not None and len(row) > stock_col:
                str_stk = str(row[stock_col]).strip()
                if str_stk != "":
                    try:
                        val_stock = max(0, int(float(str_stk)))
                    except Exception:
                        pass

            val_menudeo = None
            if men_col is not None and len(row) > men_col:
                str_m = str(row[men_col]).strip().replace("$", "").replace(",", "")
                if str_m != "":
                    try:
                        val_menudeo = float(str_m)
                    except Exception:
                        pass

            val_mayoreo = None
            if may_col is not None and len(row) > may_col:
                str_my = str(row[may_col]).strip().replace("$", "").replace(",", "")
                if str_my != "":
                    try:
                        val_mayoreo = float(str_my)
                    except Exception:
                        pass

            val_caja = None
            if caj_col is not None and len(row) > caj_col:
                str_cj = str(row[caj_col]).strip().replace("$", "").replace(",", "")
                if str_cj != "":
                    try:
                        val_caja = float(str_cj)
                    except Exception:
                        pass

            nom_item = str(row[nom_col]).strip() if nom_col is not None and len(row) > nom_col else ""
            color_item = str(row[col_col]).strip() if col_col is not None and len(row) > col_col else ""
            cat_item = str(row[cat_col]).strip() if cat_col is not None and len(row) > cat_col else ""

            if acc_id is not None:
                item_dict = {"id": acc_id}
                if val_stock is not None:
                    item_dict["stock"] = val_stock
                if val_menudeo is not None:
                    item_dict["precio_venta"] = val_menudeo
                if val_mayoreo is not None:
                    item_dict["precio_mayoreo"] = val_mayoreo
                    item_dict["precio_20_pzs"] = val_mayoreo
                if val_caja is not None:
                    item_dict["precio_caja"] = val_caja
                items_to_update.append(item_dict)
            elif nom_item:
                items_to_create.append({
                    "nombre": nom_item,
                    "color": color_item,
                    "categoria": cat_item or "General",
                    "stock": val_stock or 0,
                    "precio_venta": val_menudeo or val_mayoreo or 0.0,
                    "precio_mayoreo": val_mayoreo or val_menudeo or 0.0,
                    "precio_20_pzs": val_mayoreo or val_menudeo or 0.0,
                    "precio_caja": val_caja or val_mayoreo or 0.0
                })
    else:
        wb = openpyxl.load_workbook(io.BytesIO(contenido_bytes), data_only=True)
        ws = wb.active

        headers = [cell.value for cell in ws[1] if cell.value is not None]
        norm_headers = [str(h).strip().lower().replace(" ", "").replace("_", "") for h in headers]

        id_col = None
        stock_col = None
        nom_col = None
        col_col = None
        cat_col = None
        men_col = None
        may_col = None
        caj_col = None

        for idx, h in enumerate(norm_headers):
            if h in ("id", "codigo", "clave"):
                id_col = idx
            elif h in ("nuevostock", "stock", "stockfisico", "piezas", "cantidad"):
                stock_col = idx
            elif h in ("nombredelaccesorio", "nombre", "accesorio", "descripcion", "item"):
                nom_col = idx
            elif h in ("color", "variante"):
                col_col = idx
            elif h in ("categoria", "familia", "tipo"):
                cat_col = idx
            elif h in ("menudeo", "preciomenudeo", "precioventa", "precio"):
                men_col = idx
            elif h in ("mayoreo", "preciomayoreo"):
                may_col = idx
            elif h in ("caja", "preciocaja"):
                caj_col = idx

        if id_col is None and stock_col is None and nom_col is None:
            raise ValueError("El archivo Excel debe contener al menos las columnas 'ID' y 'NUEVO STOCK' o 'Nombre del Accesorio'.")

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row:
                continue
            acc_id = None
            if id_col is not None and len(row) > id_col and row[id_col] is not None:
                try:
                    raw_id = int(float(str(row[id_col]).strip()))
                    if raw_id > 0:
                        acc_id = raw_id
                except Exception:
                    acc_id = None

            val_stock = None
            if stock_col is not None and len(row) > stock_col and row[stock_col] is not None:
                try:
                    val_stock = max(0, int(float(str(row[stock_col]).strip())))
                except Exception:
                    pass

            val_menudeo = None
            if men_col is not None and len(row) > men_col and row[men_col] is not None:
                try:
                    val_menudeo = float(str(row[men_col]).strip().replace("$", "").replace(",", ""))
                except Exception:
                    pass

            val_mayoreo = None
            if may_col is not None and len(row) > may_col and row[may_col] is not None:
                try:
                    val_mayoreo = float(str(row[may_col]).strip().replace("$", "").replace(",", ""))
                except Exception:
                    pass

            val_caja = None
            if caj_col is not None and len(row) > caj_col and row[caj_col] is not None:
                try:
                    val_caja = float(str(row[caj_col]).strip().replace("$", "").replace(",", ""))
                except Exception:
                    pass

            nom_item = str(row[nom_col]).strip() if nom_col is not None and len(row) > nom_col and row[nom_col] is not None else ""
            color_item = str(row[col_col]).strip() if col_col is not None and len(row) > col_col and row[col_col] is not None else ""
            cat_item = str(row[cat_col]).strip() if cat_col is not None and len(row) > cat_col and row[cat_col] is not None else ""

            if acc_id is not None:
                item_dict = {"id": acc_id}
                if val_stock is not None:
                    item_dict["stock"] = val_stock
                if val_menudeo is not None:
                    item_dict["precio_venta"] = val_menudeo
                if val_mayoreo is not None:
                    item_dict["precio_mayoreo"] = val_mayoreo
                    item_dict["precio_20_pzs"] = val_mayoreo
                if val_caja is not None:
                    item_dict["precio_caja"] = val_caja
                items_to_update.append(item_dict)
            elif nom_item:
                items_to_create.append({
                    "nombre": nom_item,
                    "color": color_item,
                    "categoria": cat_item or "General",
                    "stock": val_stock or 0,
                    "precio_venta": val_menudeo or val_mayoreo or 0.0,
                    "precio_mayoreo": val_mayoreo or val_menudeo or 0.0,
                    "precio_20_pzs": val_mayoreo or val_menudeo or 0.0,
                    "precio_caja": val_caja or val_mayoreo or 0.0
                })

    if not items_to_update and not items_to_create:
        raise ValueError("No se encontraron filas con datos válidos para procesar.")

    total_creados = 0
    if items_to_create:
        for n_it in items_to_create:
            try:
                p_men = n_it.get("precio_venta") or 0.0
                p_may = n_it.get("precio_mayoreo") or p_men
                p_caj = n_it.get("precio_caja") or p_may
                p_stk = n_it.get("stock") or 0
                crud.create_accesorio(
                    nombre=n_it["nombre"],
                    categoria=n_it.get("categoria") or "General",
                    costo_compra=p_caj or p_may or p_men,
                    precio_venta=p_men or p_may,
                    precio_mayoreo=p_may,
                    precio_caja=p_caj,
                    stock=p_stk,
                    color=n_it.get("color") or "",
                    tenant_id=tenant_id
                )
                total_creados += 1
            except Exception as ex:
                errores.append(f"Error creando '{n_it.get('nombre')}': {str(ex)}")

    res = {"total_actualizados": 0, "total_creados": total_creados, "mensaje": ""}
    if items_to_update:
        res_upd = crud.actualizar_stock_masivo(
            items=items_to_update,
            motivo_general="Importación masiva desde Google Sheets / Excel",
            usuario=usuario,
            tenant_id=tenant_id
        )
        res["total_actualizados"] = res_upd.get("total_actualizados", 0)

    mensajes = []
    if res["total_actualizados"] > 0:
        mensajes.append(f"Se actualizaron {res['total_actualizados']} accesorios existentes")
    if total_creados > 0:
        mensajes.append(f"Se crearon {total_creados} nuevos accesorios")
    res["mensaje"] = ". ".join(mensajes) + "." if mensajes else "Procesado sin cambios."
    res["errores_filas"] = errores
    return res

